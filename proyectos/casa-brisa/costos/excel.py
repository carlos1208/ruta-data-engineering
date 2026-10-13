"""Exporta el presupuesto y el cronograma a una hoja de cálculo editable.

Las cantidades y los precios unitarios son valores; los totales, los indirectos y el
Gantt son fórmulas. Si se cambia una cantidad, un precio, un porcentaje o la semana de
inicio de una actividad, el libro se recalcula solo.
"""
import json
from datetime import date, timedelta
from pathlib import Path
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

AQUI = Path(__file__).parent
P = json.loads((AQUI / "presupuesto.json").read_text())
C = json.loads((AQUI / "cronograma.json").read_text())
PR = P["precios"]
ESC = ("bajo", "medio", "alto")

COP = '"$"#,##0'
MILL = '"$"#,##0.0,,"M"'
PCT = '0.0%'
B = Font(bold=True)
HEAD = PatternFill("solid", fgColor="1F2A2E")
HFONT = Font(bold=True, color="FFFFFF")
CAP = PatternFill("solid", fgColor="EEF1F2")
EDIT = PatternFill("solid", fgColor="FFF6DC")   # celdas para editar
thin = Side(style="thin", color="D5DADC")
FASE = {"pre": "2A78D6", "gris": "EB6834", "acab": "1BAF7A", "ext": "EDA100"}

def header(ws, row, names, widths=None):
    for i, n in enumerate(names, 1):
        c = ws.cell(row, i, n); c.font = HFONT; c.fill = HEAD; c.alignment = Alignment(wrap_text=True, vertical="center")
    if widths:
        for i, w in enumerate(widths, 1): ws.column_dimensions[L(i)].width = w
    ws.freeze_panes = ws.cell(row + 1, 3)

wb = Workbook()

# ------------------------------------------------------------------ Partidas
wp = wb.active; wp.title = "Partidas"
wp["A1"] = "Presupuesto de obra de Casa Brisa · costo directo por partida (COP, octubre de 2026)"; wp["A1"].font = Font(bold=True, size=13)
wp["A2"] = "Celdas amarillas: cantidades y precios unitarios editables. Totales y mano de obra se calculan."
cols = ["Capítulo", "Partida", "Unidad", "Cantidad", "PU bajo", "PU medio", "PU alto",
        "% MO bajo", "% MO medio", "% MO alto", "Total bajo", "Total medio", "Total alto",
        "MO bajo", "MO medio", "MO alto", "Base del precio", "Medición"]
header(wp, 4, cols, [26, 46, 8, 10, 13, 13, 13, 8, 8, 8, 15, 15, 15, 14, 14, 14, 60, 60])
r = 5; first = r
for p in P["partidas"]:
    wp.cell(r, 1, p["capitulo"]); wp.cell(r, 2, p["partida"]); wp.cell(r, 3, p["unidad"])
    wp.cell(r, 4, p["cantidad"]).fill = EDIT
    for j, s in enumerate(ESC):
        c = wp.cell(r, 5 + j, p[f"pu_{s}"]); c.number_format = COP; c.fill = EDIT
        c = wp.cell(r, 8 + j, round(p[f"mo_{s}"] / p[f"total_{s}"], 4) if p[f"total_{s}"] else 0); c.number_format = PCT
        c = wp.cell(r, 11 + j, f"=ROUND(D{r}*{L(5 + j)}{r},0)"); c.number_format = COP
        c = wp.cell(r, 14 + j, f"=ROUND({L(11 + j)}{r}*{L(8 + j)}{r},0)"); c.number_format = COP
    wp.cell(r, 17, p["base"]); wp.cell(r, 18, p.get("medicion", ""))
    r += 1
last = r - 1
wp.cell(r, 2, "Costo directo total").font = B
for j in range(6):
    c = wp.cell(r, 11 + j, f"=SUM({L(11 + j)}{first}:{L(11 + j)}{last})"); c.number_format = COP; c.font = B
TOTROW = r
wp.auto_filter.ref = f"A4:{L(len(cols))}{last}"

# ------------------------------------------------------------------ Capítulos
wc = wb.create_sheet("Capítulos")
header(wc, 1, ["Capítulo", "", "Total bajo", "Total medio", "Total alto", "MO medio", "% MO medio", "% del costo directo (medio)"], [40, 2, 16, 16, 16, 16, 10, 12])
caps = [c["capitulo"] for c in P["capitulos"]]
for i, cap in enumerate(caps):
    rr = i + 2; wc.cell(rr, 1, cap)
    for j, col in enumerate(("K", "L", "M")):
        c = wc.cell(rr, 3 + j, f'=SUMIF(Partidas!$A${first}:$A${last},$A{rr},Partidas!{col}${first}:{col}${last})'); c.number_format = COP
    c = wc.cell(rr, 6, f'=SUMIF(Partidas!$A${first}:$A${last},$A{rr},Partidas!O${first}:O${last})'); c.number_format = COP
    c = wc.cell(rr, 7, f"=F{rr}/D{rr}"); c.number_format = PCT
    c = wc.cell(rr, 8, f"=D{rr}/Partidas!$L${TOTROW}"); c.number_format = PCT
rr = len(caps) + 2; wc.cell(rr, 1, "Total").font = B
for j in range(4):
    c = wc.cell(rr, 3 + j, f"=SUM({L(3 + j)}2:{L(3 + j)}{rr - 1})"); c.number_format = COP; c.font = B
wc.freeze_panes = "B2"

# ------------------------------------------------------------------ Resumen con parámetros
wr = wb.create_sheet("Resumen", 0)
wr["A1"] = "Casa Brisa · costo total del proyecto"; wr["A1"].font = Font(bold=True, size=14)
wr["A2"] = "Barranquilla, precios de octubre de 2026. Celdas amarillas: parámetros editables."
for i, w in enumerate([52, 16, 16, 16, 44], 1): wr.column_dimensions[L(i)].width = w
for j, n in enumerate(["Parámetro", "Bajo", "Medio", "Alto", "Fuente"]):
    c = wr.cell(4, 1 + j, n); c.font = HFONT; c.fill = HEAD
par = [("Área construida (m²)", "area", None), ("Diseño arquitectónico (COP/m²)", "S2", COP), ("Diseño estructural (COP/m²)", "S3", COP),
       ("Diseños eléctrico e hidrosanitario (COP)", "S4", COP), ("Estudio de suelos (COP)", "S1", COP), ("Levantamiento topográfico (COP)", "S7", COP),
       ("Expensas de curaduría (COP)", "S5", COP), ("AIU (% del costo directo)", "B4", PCT), ("Interventoría (% del costo directo)", "S6", PCT),
       ("Delineación urbana (% de la obra)", "B6", PCT), ("Base mínima de delineación (COP/m²)", "B7", COP), ("Utilidad dentro del AIU", "util", PCT),
       ("IVA sobre la utilidad", "B8", PCT), ("Pólizas (% del contrato)", "S8", PCT)]
ROW = {}
for i, (n, k, fmt) in enumerate(par):
    rr = 5 + i; ROW[k] = rr; wr.cell(rr, 1, n)
    for j, s in enumerate(ESC):
        if k == "area": val = P["area"]
        elif k == "util": val = 0.30
        else: val = PR[k][s] / 100 if fmt == PCT else PR[k][s]
        c = wr.cell(rr, 2 + j, val); c.fill = EDIT
        if fmt: c.number_format = fmt
    wr.cell(rr, 5, "supuesto propio" if k in ("area", "util") else f"{k}: {PR[k]['concepto'][:60]}")
r0 = 5 + len(par) + 1
for j, n in enumerate(["Concepto", "Bajo", "Medio", "Alto", "Cómo se calcula"]):
    c = wr.cell(r0, 1 + j, n); c.font = HFONT; c.fill = HEAD
lines = [
    ("Costo directo de obra", lambda x: f"=Partidas!{ {'B':'K','C':'L','D':'M'}[x] }{TOTROW}", "suma de las partidas"),
    ("Diseño arquitectónico", lambda x: f"={x}{ROW['S2']}*{x}{ROW['area']}", "S2 × área"),
    ("Diseño estructural NSR-10", lambda x: f"={x}{ROW['S3']}*{x}{ROW['area']}", "S3 × área"),
    ("Diseños eléctrico RETIE e hidrosanitario", lambda x: f"={x}{ROW['S4']}", "S4"),
    ("Estudio de suelos", lambda x: f"={x}{ROW['S1']}", "S1"),
    ("Levantamiento topográfico", lambda x: f"={x}{ROW['S7']}", "S7"),
    ("Expensas de licencia (curaduría)", lambda x: f"={x}{ROW['S5']}", "S5"),
    ("Impuesto de delineación urbana", lambda x: f"={x}{ROW['B6']}*MAX({x}{r0+1}*(1+{x}{ROW['B4']}),{x}{ROW['B7']}*{x}{ROW['area']})", "B6 × mayor entre obra con AIU y B7 × área"),
    ("Interventoría", lambda x: f"={x}{ROW['S6']}*{x}{r0+1}", "S6 × costo directo"),
    ("AIU del constructor", lambda x: f"={x}{ROW['B4']}*{x}{r0+1}", "B4 × costo directo"),
    ("IVA sobre la utilidad", lambda x: f"={x}{ROW['B8']}*{x}{ROW['util']}*{x}{ROW['B4']}*{x}{r0+1}", "B8 × utilidad (30 % del AIU)"),
    ("Pólizas", lambda x: f"={x}{ROW['S8']}*{x}{r0+1}*(1+{x}{ROW['B4']})", "S8 × valor del contrato"),
]
for i, (n, f, how) in enumerate(lines):
    rr = r0 + 1 + i; wr.cell(rr, 1, n)
    for j, x in enumerate("BCD"):
        c = wr.cell(rr, 2 + j, f(x)); c.number_format = COP
    wr.cell(rr, 5, how)
rt = r0 + 1 + len(lines)
wr.cell(rt, 1, "Total del proyecto").font = Font(bold=True, size=12)
wr.cell(rt + 1, 1, "Total por m² construido"); wr.cell(rt + 2, 1, "Mano de obra (% del costo directo)")
for j, x in enumerate("BCD"):
    c = wr.cell(rt, 2 + j, f"=SUM({x}{r0+1}:{x}{rt-1})"); c.number_format = COP; c.font = Font(bold=True, size=12)
    c = wr.cell(rt + 1, 2 + j, f"={x}{rt}/{x}{ROW['area']}"); c.number_format = COP
    mo = {"B": "N", "C": "O", "D": "P"}[x]; tt = {"B": "K", "C": "L", "D": "M"}[x]
    c = wr.cell(rt + 2, 2 + j, f"=Partidas!{mo}{TOTROW}/Partidas!{tt}{TOTROW}"); c.number_format = PCT
wr.cell(rt + 4, 1, "Valores de referencia calculados por presupuesto.py (deben coincidir sin cambios):").font = Font(italic=True)
for j, s in enumerate(ESC):
    c = wr.cell(rt + 4, 2 + j, P["resumen"][s]["total"]); c.number_format = COP; c.font = Font(italic=True)

# ------------------------------------------------------------------ Gantt
wg = wb.create_sheet("Cronograma")
ini = date.fromisoformat(C["inicio"]); W = C["semanas"]
wg["A1"] = f"Cronograma de Casa Brisa · semana 0 = {ini.isoformat()} · {W} semanas"; wg["A1"].font = Font(bold=True, size=13)
wg["A2"] = "Edite la semana de inicio o la duración (amarillo): las barras se redibujan. Las fechas no reprograman las sucesoras; para eso, cronograma.py."
hd = ["ID", "Actividad", "Fase", "Inicio (sem)", "Duración (sem)", "Fecha inicio", "Fecha fin", "Holgura (sem)", "Ruta crítica"]
for i, n in enumerate(hd, 1):
    c = wg.cell(4, i, n); c.font = HFONT; c.fill = HEAD; c.alignment = Alignment(wrap_text=True)
for i, w in enumerate([5, 44, 10, 8, 9, 11, 11, 8, 8], 1): wg.column_dimensions[L(i)].width = w
G0 = len(hd) + 1
for w in range(W):
    col = G0 + w; wg.column_dimensions[L(col)].width = 2.6
    d = ini + timedelta(weeks=w)
    c = wg.cell(3, col, d.strftime("%b %y") if d.day <= 7 else None); c.font = Font(size=8)
    c = wg.cell(4, col, w); c.font = Font(size=7, color="FFFFFF"); c.fill = HEAD
for i, a in enumerate(C["actividades"]):
    rr = 5 + i
    wg.cell(rr, 1, a["id"]); wg.cell(rr, 2, a["nombre"]).font = B if a["critica"] else Font()
    wg.cell(rr, 3, C["fases"][a["fase"]]); wg.cell(rr, 4, a["es"]).fill = EDIT; wg.cell(rr, 5, a["dur"]).fill = EDIT
    c = wg.cell(rr, 6, f'=DATE({ini.year},{ini.month},{ini.day})+7*D{rr}'); c.number_format = "dd-mmm-yy"
    c = wg.cell(rr, 7, f"=F{rr}+7*E{rr}-3"); c.number_format = "dd-mmm-yy"
    wg.cell(rr, 8, a["holgura"]); wg.cell(rr, 9, "sí" if a["critica"] else "")
    for w in range(W): wg.cell(rr, G0 + w).border = Border(left=thin) if (ini + timedelta(weeks=w)).day <= 7 else Border()
lastA = 4 + len(C["actividades"])
rng = f"{L(G0)}5:{L(G0 + W - 1)}{lastA}"
for k, col in FASE.items():
    wg.conditional_formatting.add(rng, FormulaRule(
        formula=[f'AND({L(G0)}$4>=$D5,{L(G0)}$4<$D5+$E5,$C5="{C["fases"][k]}")'], fill=PatternFill("solid", fgColor=col, bgColor=col)))
rh = lastA + 2; wg.cell(rh, 2, "Hitos").font = B
for i, h in enumerate(C["hitos"]):
    wg.cell(rh + 1 + i, 2, h["nombre"]); c = wg.cell(rh + 1 + i, 6, date.fromisoformat(h["fecha"])); c.number_format = "dd-mmm-yy"
    wg.cell(rh + 1 + i, G0 + h["semana"] - 1, "◆").alignment = Alignment(horizontal="center")
wg.freeze_panes = wg.cell(5, G0)

# ------------------------------------------------------------------ Flujo de caja
wf = wb.create_sheet("Flujo de caja")
for i, n in enumerate(["Mes", "Bajo", "Medio", "Alto", "Acumulado medio", "% acumulado medio"], 1):
    c = wf.cell(1, i, n); c.font = HFONT; c.fill = HEAD; wf.column_dimensions[L(i)].width = 16
n = len(C["flujo_medio"])
for i, m in enumerate(C["flujo_medio"]):
    rr = i + 2; wf.cell(rr, 1, m["mes"])
    for j, s in enumerate(ESC):
        c = wf.cell(rr, 2 + j, C[f"flujo_{s}"][i]["valor"]); c.number_format = COP
    c = wf.cell(rr, 5, f"=SUM($C$2:C{rr})"); c.number_format = COP
    c = wf.cell(rr, 6, f"=E{rr}/SUM($C$2:$C${n + 1})"); c.number_format = PCT
wf.cell(n + 2, 1, "Total").font = B
for j in range(3):
    c = wf.cell(n + 2, 2 + j, f"=SUM({L(2 + j)}2:{L(2 + j)}{n + 1})"); c.number_format = COP; c.font = B

# ------------------------------------------------------------------ Precios y fuentes
wq = wb.create_sheet("Precios y fuentes")
header(wq, 1, ["Código", "Concepto", "Unidad", "Bajo", "Medio", "Alto", "Confianza", "Fuente principal", "Enlace", "Nota"], [8, 50, 18, 13, 13, 13, 10, 36, 50, 80])
for i, p in enumerate(PR.values()):
    rr = i + 2; f0 = (p.get("fuentes") or [{}])[0]
    vals = [p["id"], p["concepto"], p["unidad"], p["bajo"], p["medio"], p["alto"], p["confianza"], f0.get("nombre", ""), f0.get("url", ""), p.get("nota", "")]
    for j, v in enumerate(vals):
        c = wq.cell(rr, 1 + j, v)
        if 3 <= j <= 5 and isinstance(v, (int, float)) and v > 100: c.number_format = COP
    if f0.get("url"): wq.cell(rr, 9).hyperlink = f0["url"]
wq.cell(len(PR) + 3, 2, "Supuestos propios").font = B
for i, s in enumerate(P["supuestos"]):
    rr = len(PR) + 4 + i; wq.cell(rr, 2, s["texto"])
    for j, e in enumerate(ESC): c = wq.cell(rr, 4 + j, s[e]); c.number_format = COP

out = AQUI / "presupuesto-casa-brisa.xlsx"
wb.save(out)
print("ok", out.name)
